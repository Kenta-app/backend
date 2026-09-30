"""Validate two blind annotation workbooks and create an adjudication log."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


VALID_LABELS = {"agree", "disagree", "discuss", "unrelated"}
VALID_CONFIDENCE = {"alta", "media", "baja"}
VALID_QUALITY = {"include", "review", "exclude"}
IMMUTABLE_FIELDS = ("pair_id", "source_name", "claim", "article", "original_url", "article_id", "group_id")


def normalize(value: object) -> str:
    return str(value or "").strip()


def read_blind_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_annotation_workbook(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if "Anotación" not in workbook.sheetnames:
        raise ValueError(f"{path.name}: missing Anotación sheet")
    sheet = workbook["Anotación"]
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize(value) for value in next(rows)]
    required = set(IMMUTABLE_FIELDS) | {"evidence_span", "label", "confidence", "notes", "quality_flag", "annotator"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"{path.name}: missing columns {missing}")
    output = []
    for values in rows:
        row = {header: normalize(values[index]) for index, header in enumerate(headers)}
        if row.get("pair_id"):
            output.append(row)
    return output


def validate_annotations(
    blind: list[dict[str, str]], annotations: list[dict[str, str]], workbook_name: str
) -> dict[str, dict[str, str]]:
    blind_by_id = {row["pair_id"]: row for row in blind}
    annotation_by_id = {row["pair_id"]: row for row in annotations}
    if len(annotation_by_id) != len(annotations):
        raise ValueError(f"{workbook_name}: duplicate pair_id")
    if set(blind_by_id) != set(annotation_by_id):
        missing = sorted(set(blind_by_id) - set(annotation_by_id))
        extra = sorted(set(annotation_by_id) - set(blind_by_id))
        raise ValueError(f"{workbook_name}: pair mismatch; missing={missing[:5]}, extra={extra[:5]}")
    errors = []
    for pair_id, row in annotation_by_id.items():
        seed = blind_by_id[pair_id]
        for field in IMMUTABLE_FIELDS:
            if normalize(row.get(field)) != normalize(seed.get(field)):
                errors.append(f"{pair_id}: modified {field}")
        label = row["label"].lower()
        confidence = row["confidence"].lower()
        quality = row["quality_flag"].lower()
        if label not in VALID_LABELS:
            errors.append(f"{pair_id}: invalid or missing label {label!r}")
        if confidence not in VALID_CONFIDENCE:
            errors.append(f"{pair_id}: invalid or missing confidence {confidence!r}")
        if quality not in VALID_QUALITY:
            errors.append(f"{pair_id}: invalid quality_flag {quality!r}")
        if not row["evidence_span"]:
            errors.append(f"{pair_id}: missing evidence_span")
        if confidence == "baja" and not row["notes"]:
            errors.append(f"{pair_id}: notes required for baja confidence")
    if errors:
        raise ValueError(f"{workbook_name}: " + "; ".join(errors[:20]))
    return annotation_by_id


def cohen_kappa(left: list[str], right: list[str]) -> float:
    total = len(left)
    observed = sum(a == b for a, b in zip(left, right, strict=True)) / total
    left_counts, right_counts = Counter(left), Counter(right)
    expected = sum(left_counts[label] * right_counts[label] for label in VALID_LABELS) / (total * total)
    return (observed - expected) / (1 - expected) if expected < 1 else 1.0


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--blind-csv", type=Path, required=True)
    parser.add_argument("--annotator-1", type=Path, required=True)
    parser.add_argument("--annotator-2", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    blind = read_blind_csv(args.blind_csv)
    left = validate_annotations(blind, read_annotation_workbook(args.annotator_1), args.annotator_1.name)
    right = validate_annotations(blind, read_annotation_workbook(args.annotator_2), args.annotator_2.name)
    integrated = []
    for seed in blind:
        pair_id = seed["pair_id"]
        one, two = left[pair_id], right[pair_id]
        agreement = one["label"].lower() == two["label"].lower()
        integrated.append({
            **seed,
            "annotator_1_name": one["annotator"], "label_annotator_1": one["label"].lower(),
            "evidence_annotator_1": one["evidence_span"], "confidence_annotator_1": one["confidence"].lower(),
            "notes_annotator_1": one["notes"], "quality_annotator_1": one["quality_flag"].lower(),
            "annotator_2_name": two["annotator"], "label_annotator_2": two["label"].lower(),
            "evidence_annotator_2": two["evidence_span"], "confidence_annotator_2": two["confidence"].lower(),
            "notes_annotator_2": two["notes"], "quality_annotator_2": two["quality_flag"].lower(),
            "agreement_status": "agreement" if agreement else "disagreement",
            "adjudicated_label": one["label"].lower() if agreement else "",
            "adjudication_note": "Acuerdo independiente" if agreement else "",
            "final_label": one["label"].lower() if agreement else "",
        })
    annotation_fields = list(integrated[0])
    disagreements = [row for row in integrated if row["agreement_status"] == "disagreement"]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_csv(args.output_dir / "annotation_log.csv", integrated, annotation_fields)
    write_csv(args.output_dir / "disagreements_for_adjudication.csv", disagreements, annotation_fields)
    labels_one = [row["label_annotator_1"] for row in integrated]
    labels_two = [row["label_annotator_2"] for row in integrated]
    report = {
        "pairs": len(integrated), "agreements": len(integrated) - len(disagreements),
        "disagreements": len(disagreements),
        "raw_agreement": (len(integrated) - len(disagreements)) / len(integrated),
        "cohen_kappa": cohen_kappa(labels_one, labels_two),
        "annotator_1_counts": dict(Counter(labels_one)), "annotator_2_counts": dict(Counter(labels_two)),
        "ready_for_adjudication": True,
    }
    (args.output_dir / "agreement_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
