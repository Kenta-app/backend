"""Export the first worksheet of a completed annotation workbook as UTF-8 CSV."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = (
    "order",
    "evidence_id",
    "claim",
    "source_name",
    "source_title",
    "evidence_excerpt",
    "source_url",
    "review_status",
    "human_stance",
    "confidence",
    "notes",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook = load_workbook(args.input, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows))
        if headers != EXPECTED_HEADERS:
            raise SystemExit(
                f"Unexpected headers in {args.input}: {headers!r}"
            )

        exported = []
        for values in rows:
            row = {
                header: "" if value is None else str(value).strip()
                for header, value in zip(headers, values, strict=True)
            }
            if row["evidence_id"]:
                exported.append(row)
    finally:
        workbook.close()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXPECTED_HEADERS)
        writer.writeheader()
        writer.writerows(exported)

    print(f"Exported {len(exported)} rows from {args.input} to {args.output}")


if __name__ == "__main__":
    main()
