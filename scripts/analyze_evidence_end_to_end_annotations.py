"""Integrate the two blind end-to-end evidence annotation workbooks.

This analysis is deliberately pre-adjudication. Exact agreements define the
consensus subset; disagreements are exported for later human adjudication.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


STANCE_LABELS = ("supports", "refutes", "discusses", "unrelated")
COVERAGE_LABELS = ("sufficient", "partial", "none", "no_evidence_found", "unable_to_assess")
REVIEW_STATUSES = ("evaluable", "unavailable", "insufficient_context")
CONFIDENCE_LABELS = ("high", "medium", "low")
MISSING_LABELS = ("yes", "no", "unable_to_assess")
SOCIAL_CONTEXT_LABELS = ("correct", "incorrect", "unable_to_assess")


def norm(value: Any) -> str:
    return str(value or "").strip()


def read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [norm(value) for value in next(rows)]
        return [dict(zip(headers, row)) for row in rows if any(value not in (None, "") for value in row)]
    finally:
        workbook.close()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def index_unique(rows: Iterable[dict[str, Any]], key: str) -> dict[str, dict[str, Any]]:
    indexed: dict[str, dict[str, Any]] = {}
    for row in rows:
        value = norm(row.get(key))
        if not value or value in indexed:
            raise ValueError(f"Missing or duplicate {key}: {value!r}")
        indexed[value] = row
    return indexed


def validate_source_annotations(
    label: str,
    rows: dict[str, dict[str, Any]],
    system_rows: dict[str, dict[str, Any]],
) -> None:
    errors: list[str] = []
    frozen_fields = (
        "case_id", "claim", "original_source", "source_name", "source_title",
        "evidence_excerpt", "source_url",
    )
    for evidence_id, row in rows.items():
        key = system_rows.get(evidence_id)
        if key is None:
            continue
        for field in frozen_fields:
            if norm(row.get(field)) != norm(key.get(field)):
                errors.append(f"{evidence_id}: frozen field {field} was modified")
        status = norm(row.get("review_status"))
        stance = norm(row.get("human_stance"))
        confidence = norm(row.get("confidence"))
        notes = norm(row.get("notes"))
        if status not in REVIEW_STATUSES:
            errors.append(f"{evidence_id}: invalid review_status {status!r}")
        if status == "evaluable":
            if stance not in STANCE_LABELS:
                errors.append(f"{evidence_id}: evaluable row requires a valid human_stance")
            if confidence not in CONFIDENCE_LABELS:
                errors.append(f"{evidence_id}: evaluable row requires confidence")
        elif stance or confidence:
            errors.append(f"{evidence_id}: non-evaluable row must leave stance/confidence blank")
        if not notes:
            errors.append(f"{evidence_id}: notes are required")
    if errors:
        raise ValueError(f"Invalid source annotations for {label}:\n- " + "\n- ".join(errors))


def validate_coverage_annotations(
    label: str,
    rows: dict[str, dict[str, Any]],
    system_rows: dict[str, dict[str, Any]],
) -> None:
    errors: list[str] = []
    frozen_fields = (
        "platform", "claim", "publication_context", "social_context_completeness",
        "social_reference_count", "social_reference_types", "original_source", "original_url",
    )
    for case_id, row in rows.items():
        key = system_rows.get(case_id)
        if key is None:
            continue
        for field in frozen_fields:
            if norm(row.get(field)) != norm(key.get(field)):
                errors.append(f"{case_id}: frozen field {field} was modified")
        try:
            returned_count = int(row.get("returned_source_count") or 0)
        except (TypeError, ValueError):
            returned_count = -1
        if returned_count != int(key.get("returned_source_count") or 0):
            errors.append(f"{case_id}: frozen returned_source_count was modified")
        status = norm(row.get("coverage_status"))
        missing = norm(row.get("missing_relevant_evidence"))
        missing_url = norm(row.get("missing_source_url"))
        if status not in COVERAGE_LABELS:
            errors.append(f"{case_id}: invalid coverage_status {status!r}")
        if missing not in MISSING_LABELS:
            errors.append(f"{case_id}: invalid missing_relevant_evidence {missing!r}")
        if missing == "yes" and not missing_url:
            errors.append(f"{case_id}: missing_source_url is required when missing=yes")
        platform = norm(row.get("platform"))
        social_status = norm(row.get("social_context_status"))
        if platform == "twitter" and social_status not in SOCIAL_CONTEXT_LABELS:
            errors.append(f"{case_id}: twitter row requires a valid social_context_status")
        if platform != "twitter" and social_status != "not_applicable":
            errors.append(f"{case_id}: non-social row requires social_context_status=not_applicable")
        if not norm(row.get("notes")):
            errors.append(f"{case_id}: notes are required")
    if errors:
        raise ValueError(f"Invalid coverage annotations for {label}:\n- " + "\n- ".join(errors))


def cohen_kappa(left: list[str], right: list[str]) -> float | None:
    if len(left) != len(right) or not left:
        return None
    labels = sorted(set(left) | set(right))
    observed = sum(a == b for a, b in zip(left, right)) / len(left)
    left_counts = Counter(left)
    right_counts = Counter(right)
    expected = sum(
        (left_counts[label] / len(left)) * (right_counts[label] / len(right))
        for label in labels
    )
    if math.isclose(expected, 1.0):
        return 1.0 if math.isclose(observed, 1.0) else None
    return (observed - expected) / (1 - expected)


def wilson_interval(successes: int, total: int, z: float = 1.959963984540054) -> list[float] | None:
    if total == 0:
        return None
    proportion = successes / total
    denominator = 1 + z * z / total
    center = (proportion + z * z / (2 * total)) / denominator
    margin = z * math.sqrt(proportion * (1 - proportion) / total + z * z / (4 * total * total)) / denominator
    return [max(0.0, center - margin), min(1.0, center + margin)]


def classification_metrics(gold: list[str], predicted: list[str], labels: tuple[str, ...]) -> dict[str, Any]:
    if len(gold) != len(predicted):
        raise ValueError("Gold and prediction lengths differ")
    accuracy_count = sum(a == b for a, b in zip(gold, predicted))
    per_class: dict[str, dict[str, Any]] = {}
    f1_values = []
    observed_f1_values = []
    for label in labels:
        tp = sum(g == label and p == label for g, p in zip(gold, predicted))
        fp = sum(g != label and p == label for g, p in zip(gold, predicted))
        fn = sum(g == label and p != label for g, p in zip(gold, predicted))
        support = sum(g == label for g in gold)
        precision = tp / (tp + fp) if tp + fp else 0.0
        recall = tp / (tp + fn) if tp + fn else 0.0
        f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
        per_class[label] = {
            "support": support,
            "precision": precision,
            "recall": recall,
            "f1": f1,
        }
        f1_values.append(f1)
        if support:
            observed_f1_values.append(f1)
    return {
        "n": len(gold),
        "accuracy": accuracy_count / len(gold) if gold else None,
        "accuracy_count": accuracy_count,
        "accuracy_wilson_95": wilson_interval(accuracy_count, len(gold)),
        "macro_f1_fixed_4": sum(f1_values) / len(f1_values) if f1_values else None,
        "macro_f1_observed_gold_classes": (
            sum(observed_f1_values) / len(observed_f1_values) if observed_f1_values else None
        ),
        "per_class": per_class,
    }


def pct(numerator: int, denominator: int) -> float | None:
    return numerator / denominator if denominator else None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--annotator-a", type=Path, required=True)
    parser.add_argument("--annotator-b", type=Path, required=True)
    parser.add_argument("--system-key", type=Path, required=True)
    parser.add_argument("--coverage-key", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument(
        "--preregistered-v4",
        action="store_true",
        help="Apply the five preregistered v4 acceptance criteria.",
    )
    parser.add_argument(
        "--preregistered-v7",
        action="store_true",
        help="Apply the six preregistered final confirmatory v7 acceptance criteria.",
    )
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    a_sources = index_unique(read_sheet(args.annotator_a, "Fuentes"), "evidence_id")
    b_sources = index_unique(read_sheet(args.annotator_b, "Fuentes"), "evidence_id")
    system_sources = index_unique(read_csv(args.system_key), "evidence_id")
    if set(a_sources) != set(b_sources) or set(a_sources) != set(system_sources):
        raise ValueError("Source evidence ids do not match across annotators and system key")
    validate_source_annotations("annotator A", a_sources, system_sources)
    validate_source_annotations("annotator B", b_sources, system_sources)

    source_comparison: list[dict[str, Any]] = []
    source_disagreements: list[dict[str, Any]] = []
    for evidence_id in sorted(system_sources):
        key = system_sources[evidence_id]
        a = a_sources[evidence_id]
        b = b_sources[evidence_id]
        a_evaluable = norm(a.get("review_status")) == "evaluable"
        b_evaluable = norm(b.get("review_status")) == "evaluable"
        evaluability_agreement = a_evaluable == b_evaluable
        stance_agreement = a_evaluable and b_evaluable and norm(a.get("human_stance")) == norm(b.get("human_stance"))
        consensus_stance = norm(a.get("human_stance")) if stance_agreement else ""
        needs_adjudication = not (evaluability_agreement and (not a_evaluable or stance_agreement))
        row = {
            "evidence_id": evidence_id,
            "case_id": norm(a.get("case_id")),
            "claim": norm(a.get("claim")),
            "source_name": norm(a.get("source_name")),
            "source_title": norm(a.get("source_title")),
            "evidence_excerpt": norm(a.get("evidence_excerpt")),
            "source_url": norm(a.get("source_url")),
            "model_stance": norm(key.get("model_stance")),
            "annotator_a_status": norm(a.get("review_status")),
            "annotator_a_stance": norm(a.get("human_stance")),
            "annotator_a_confidence": norm(a.get("confidence")),
            "annotator_a_notes": norm(a.get("notes")),
            "annotator_b_status": norm(b.get("review_status")),
            "annotator_b_stance": norm(b.get("human_stance")),
            "annotator_b_confidence": norm(b.get("confidence")),
            "annotator_b_notes": norm(b.get("notes")),
            "evaluability_agreement": evaluability_agreement,
            "stance_agreement": stance_agreement,
            "consensus_stance": consensus_stance,
            "model_matches_consensus": bool(consensus_stance and consensus_stance == norm(key.get("model_stance"))),
            "needs_adjudication": needs_adjudication,
            "final_status": "",
            "final_stance": "",
            "adjudication_rationale": "",
        }
        source_comparison.append(row)
        if needs_adjudication:
            source_disagreements.append(row)

    a_coverage = index_unique(read_sheet(args.annotator_a, "Cobertura"), "case_id")
    b_coverage = index_unique(read_sheet(args.annotator_b, "Cobertura"), "case_id")
    coverage_key = index_unique(read_csv(args.coverage_key), "case_id")
    if set(a_coverage) != set(b_coverage) or set(a_coverage) != set(coverage_key):
        raise ValueError("Coverage case ids do not match across annotators and system key")
    validate_coverage_annotations("annotator A", a_coverage, coverage_key)
    validate_coverage_annotations("annotator B", b_coverage, coverage_key)

    coverage_comparison: list[dict[str, Any]] = []
    coverage_disagreements: list[dict[str, Any]] = []
    for case_id in sorted(coverage_key):
        key = coverage_key[case_id]
        a = a_coverage[case_id]
        b = b_coverage[case_id]
        status_agreement = norm(a.get("coverage_status")) == norm(b.get("coverage_status"))
        missing_agreement = norm(a.get("missing_relevant_evidence")) == norm(b.get("missing_relevant_evidence"))
        is_social = norm(a.get("platform")) == "twitter"
        social_agreement = (
            norm(a.get("social_context_status")) == norm(b.get("social_context_status"))
            if is_social else True
        )
        needs_adjudication = not (status_agreement and missing_agreement and social_agreement)
        row = {
            "case_id": case_id,
            "platform": norm(a.get("platform")),
            "claim": norm(a.get("claim")),
            "publication_context": norm(key.get("publication_context")),
            "social_context_completeness": norm(key.get("social_context_completeness")),
            "social_reference_count": int(key.get("social_reference_count") or 0),
            "social_reference_types": norm(key.get("social_reference_types")),
            "original_source": norm(a.get("original_source")),
            "original_url": norm(a.get("original_url")),
            "returned_source_count": int(key.get("returned_source_count") or 0),
            "annotator_a_coverage": norm(a.get("coverage_status")),
            "annotator_a_missing": norm(a.get("missing_relevant_evidence")),
            "annotator_a_missing_url": norm(a.get("missing_source_url")),
            "annotator_a_social_context": norm(a.get("social_context_status")),
            "annotator_a_notes": norm(a.get("notes")),
            "annotator_b_coverage": norm(b.get("coverage_status")),
            "annotator_b_missing": norm(b.get("missing_relevant_evidence")),
            "annotator_b_missing_url": norm(b.get("missing_source_url")),
            "annotator_b_social_context": norm(b.get("social_context_status")),
            "annotator_b_notes": norm(b.get("notes")),
            "coverage_agreement": status_agreement,
            "missing_agreement": missing_agreement,
            "social_context_agreement": social_agreement,
            "consensus_coverage": norm(a.get("coverage_status")) if status_agreement else "",
            "consensus_missing": norm(a.get("missing_relevant_evidence")) if missing_agreement else "",
            "consensus_social_context": (
                norm(a.get("social_context_status")) if is_social and social_agreement else ""
            ),
            "needs_adjudication": needs_adjudication,
            "final_coverage": "",
            "final_missing": "",
            "final_missing_url": "",
            "final_social_context": "",
            "adjudication_rationale": "",
        }
        coverage_comparison.append(row)
        if needs_adjudication:
            coverage_disagreements.append(row)

    both_evaluable = [
        row for row in source_comparison
        if row["annotator_a_status"] == "evaluable" and row["annotator_b_status"] == "evaluable"
    ]
    consensus_stance_rows = [row for row in source_comparison if row["consensus_stance"]]
    relevant_consensus = [row for row in consensus_stance_rows if row["consensus_stance"] != "unrelated"]
    stance_metrics = classification_metrics(
        [row["consensus_stance"] for row in consensus_stance_rows],
        [row["model_stance"] for row in consensus_stance_rows],
        STANCE_LABELS,
    )

    coverage_consensus = [row for row in coverage_comparison if row["consensus_coverage"]]
    evidence_exists_consensus = [
        row for row in coverage_consensus
        if row["consensus_coverage"] in {"sufficient", "partial", "none"}
    ]
    useful_coverage = [
        row for row in evidence_exists_consensus
        if row["consensus_coverage"] in {"sufficient", "partial"}
    ]
    sufficient_coverage = [
        row for row in evidence_exists_consensus if row["consensus_coverage"] == "sufficient"
    ]
    zero_source_consensus = [
        row for row in coverage_consensus if row["returned_source_count"] == 0
    ]
    correct_empty = [
        row for row in zero_source_consensus if row["consensus_coverage"] == "no_evidence_found"
    ]
    false_empty = [row for row in zero_source_consensus if row["consensus_coverage"] == "none"]
    false_empty_rate = pct(len(false_empty), len(zero_source_consensus))
    social_cases = [row for row in coverage_comparison if row["platform"] == "twitter"]
    social_consensus = [row for row in social_cases if row["consensus_social_context"]]
    social_correct = [row for row in social_consensus if row["consensus_social_context"] == "correct"]

    a_eval_binary = ["evaluable" if row["annotator_a_status"] == "evaluable" else "not_evaluable" for row in source_comparison]
    b_eval_binary = ["evaluable" if row["annotator_b_status"] == "evaluable" else "not_evaluable" for row in source_comparison]
    summary = {
        "analysis_stage": "pre_adjudication",
        "source_pairs": len(source_comparison),
        "coverage_cases": len(coverage_comparison),
        "source_agreement": {
            "evaluability_exact": pct(sum(a == b for a, b in zip(a_eval_binary, b_eval_binary)), len(source_comparison)),
            "evaluability_kappa": cohen_kappa(a_eval_binary, b_eval_binary),
            "both_evaluable_n": len(both_evaluable),
            "stance_exact_among_both_evaluable": pct(
                sum(row["annotator_a_stance"] == row["annotator_b_stance"] for row in both_evaluable),
                len(both_evaluable),
            ),
            "stance_kappa_among_both_evaluable": cohen_kappa(
                [row["annotator_a_stance"] for row in both_evaluable],
                [row["annotator_b_stance"] for row in both_evaluable],
            ),
            "consensus_stance_n": len(consensus_stance_rows),
            "needs_adjudication_n": len(source_disagreements),
        },
        "retrieval_source_quality_consensus_only": {
            "relevant_n": len(relevant_consensus),
            "evaluable_consensus_n": len(consensus_stance_rows),
            "precision": pct(len(relevant_consensus), len(consensus_stance_rows)),
            "precision_wilson_95": wilson_interval(len(relevant_consensus), len(consensus_stance_rows)),
        },
        "system_stance_consensus_only": stance_metrics,
        "coverage_agreement": {
            "coverage_exact": pct(sum(row["coverage_agreement"] for row in coverage_comparison), len(coverage_comparison)),
            "coverage_kappa": cohen_kappa(
                [row["annotator_a_coverage"] for row in coverage_comparison],
                [row["annotator_b_coverage"] for row in coverage_comparison],
            ),
            "missing_exact": pct(sum(row["missing_agreement"] for row in coverage_comparison), len(coverage_comparison)),
            "missing_kappa": cohen_kappa(
                [row["annotator_a_missing"] for row in coverage_comparison],
                [row["annotator_b_missing"] for row in coverage_comparison],
            ),
            "consensus_coverage_n": len(coverage_consensus),
            "needs_adjudication_n": len(coverage_disagreements),
        },
        "social_context_agreement": {
            "social_cases_n": len(social_cases),
            "exact": pct(sum(row["social_context_agreement"] for row in social_cases), len(social_cases)),
            "kappa": cohen_kappa(
                [row["annotator_a_social_context"] for row in social_cases],
                [row["annotator_b_social_context"] for row in social_cases],
            ),
            "consensus_n": len(social_consensus),
            "correct_n": len(social_correct),
            "correct_rate": pct(len(social_correct), len(social_consensus)),
        },
        "coverage_consensus_only": {
            "status_counts": dict(Counter(row["consensus_coverage"] for row in coverage_consensus)),
            "evidence_exists_n": len(evidence_exists_consensus),
            "useful_coverage_n": len(useful_coverage),
            "useful_coverage_rate": pct(len(useful_coverage), len(evidence_exists_consensus)),
            "sufficient_coverage_n": len(sufficient_coverage),
            "sufficient_coverage_rate": pct(len(sufficient_coverage), len(evidence_exists_consensus)),
            "zero_source_consensus_n": len(zero_source_consensus),
            "correct_empty_n": len(correct_empty),
            "false_empty_n": len(false_empty),
            "false_empty_rate": false_empty_rate,
            "consensus_missing_yes_n": sum(row["consensus_missing"] == "yes" for row in coverage_comparison),
        },
        "system_run": {
            "retrieval_success_n": sum(row["returned_source_count"] > 0 for row in coverage_comparison),
            "retrieval_success_rate": pct(sum(row["returned_source_count"] > 0 for row in coverage_comparison), len(coverage_comparison)),
            "retrieval_success_wilson_95": wilson_interval(
                sum(row["returned_source_count"] > 0 for row in coverage_comparison),
                len(coverage_comparison),
            ),
        },
        "limitations": [
            "All quality metrics are pre-adjudication and use only exact-consensus rows.",
            "The sample is a small stratified pilot and does not estimate natural feed prevalence.",
            "Social-context quality is measured only on the six preregistered X cases.",
            "Macro-F1 over four fixed classes is depressed when a gold class is absent; observed-class macro-F1 is also reported.",
        ],
    }
    if args.preregistered_v7:
        criteria = {
            "retrieval_content_at_least_80pct": {
                "value": summary["system_run"]["retrieval_success_rate"],
                "threshold": 0.80,
                "passes": summary["system_run"]["retrieval_success_rate"] >= 0.80,
                "status": "final",
            },
            "source_precision_at_least_90pct": {
                "value": summary["retrieval_source_quality_consensus_only"]["precision"],
                "threshold": 0.90,
                "passes": summary["retrieval_source_quality_consensus_only"]["precision"] >= 0.90,
                "status": "pre_adjudication",
            },
            "useful_coverage_at_least_80pct": {
                "value": summary["coverage_consensus_only"]["useful_coverage_rate"],
                "threshold": 0.80,
                "passes": summary["coverage_consensus_only"]["useful_coverage_rate"] >= 0.80,
                "status": "pre_adjudication",
            },
            "stance_accuracy_at_least_75pct": {
                "value": summary["system_stance_consensus_only"]["accuracy"],
                "threshold": 0.75,
                "passes": summary["system_stance_consensus_only"]["accuracy"] >= 0.75,
                "status": "pre_adjudication",
            },
            "false_empty_rate_at_most_20pct": {
                "value": false_empty_rate,
                "threshold": 0.20,
                "passes": false_empty_rate is not None and false_empty_rate <= 0.20,
                "status": "pre_adjudication",
            },
            "social_context_correct_at_least_80pct": {
                "value": summary["social_context_agreement"]["correct_rate"],
                "threshold": 0.80,
                "passes": summary["social_context_agreement"]["correct_rate"] is not None
                and summary["social_context_agreement"]["correct_rate"] >= 0.80,
                "status": "pre_adjudication",
            },
        }
        summary["preregistered_criteria"] = criteria
        summary["overall_status"] = (
            "failed_retrieval_criterion"
            if not criteria["retrieval_content_at_least_80pct"]["passes"]
            else "pending_adjudication"
        )
        summary["limitations"].append(
            "The six v7 thresholds were preregistered; human-dependent results remain preliminary until adjudication."
        )
    elif args.preregistered_v4:
        criteria = {
            "retrieval_16_of_20": {
                "value": summary["system_run"]["retrieval_success_n"],
                "threshold": 16,
                "passes": summary["system_run"]["retrieval_success_n"] >= 16,
                "status": "final",
            },
            "source_precision_at_least_90pct": {
                "value": summary["retrieval_source_quality_consensus_only"]["precision"],
                "threshold": 0.90,
                "passes": summary["retrieval_source_quality_consensus_only"]["precision"] >= 0.90,
                "status": "pre_adjudication",
            },
            "useful_coverage_at_least_80pct": {
                "value": summary["coverage_consensus_only"]["useful_coverage_rate"],
                "threshold": 0.80,
                "passes": summary["coverage_consensus_only"]["useful_coverage_rate"] >= 0.80,
                "status": "pre_adjudication",
            },
            "stance_accuracy_at_least_75pct": {
                "value": summary["system_stance_consensus_only"]["accuracy"],
                "threshold": 0.75,
                "passes": summary["system_stance_consensus_only"]["accuracy"] >= 0.75,
                "status": "pre_adjudication",
            },
            "false_empty_rate_at_most_20pct": {
                "value": false_empty_rate,
                "threshold": 0.20,
                "passes": false_empty_rate is not None and false_empty_rate <= 0.20,
                "status": "pre_adjudication",
            },
        }
        summary["preregistered_criteria"] = criteria
        summary["overall_status"] = (
            "failed_retrieval_criterion"
            if not criteria["retrieval_16_of_20"]["passes"]
            else "pending_adjudication"
        )
        summary["limitations"].append(
            "Acceptance thresholds were preregistered, but source quality, coverage, stance, and false-empty results remain preliminary until adjudication."
        )
    else:
        summary["limitations"].append(
            "No pass/fail threshold was preregistered before this run."
        )

    source_fields = list(source_comparison[0])
    coverage_fields = list(coverage_comparison[0])
    write_csv(args.output_dir / "source_comparison.csv", source_comparison, source_fields)
    write_csv(args.output_dir / "source_disagreements.csv", source_disagreements, source_fields)
    write_csv(args.output_dir / "coverage_comparison.csv", coverage_comparison, coverage_fields)
    write_csv(args.output_dir / "coverage_disagreements.csv", coverage_disagreements, coverage_fields)
    (args.output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
